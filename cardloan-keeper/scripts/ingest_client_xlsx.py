#!/usr/bin/env python3
"""クライアント公式レギュレーションxlsxの一括取込

カードローン案件_レギュレーション_YUKATAN様.xlsx（ASP提供・複数社シート）を
regulations/raw/<slug>/wb_<sheet>.csv に正規化ダンプする。
チャンク化は regstore.py の CSV パーサに一本化（xlsx直パースを二重実装しない）。

実行: ./.venv/bin/python scripts/ingest_client_xlsx.py [xlsxパス]
再実行時は各 slug の wb_*.csv を消してから書き直す（冪等）。
原本xlsxは regulations/raw/_workbook/ に日付付きで保全（gitignore対象）。
"""
from __future__ import annotations

import csv
import shutil
import sys
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "regulations" / "raw"
DEFAULT_XLSX = Path.home() / "Desktop" / "カードローン案件_レギュレーション_YUKATAN様.xlsx"

# シート名 → 保存先slug。除外KW・スペックリストも含め全シートを取り込む
SHEET_MAP = {
    "案件複合_レギュレーション": "fukugou",
    "プロミスレギュレーション": "promise",
    "プロミス_除外KW": "promise",
    "モビットレギュレーション": "smbc-mobit",
    "モビット_除外KW": "smbc-mobit",
    "アイフルレギュレーション": "aiful",
    "アイフルAGレギュレーション": "aiful",
    "アイフル 商品スペックリスト": "aiful",
    "アイフル新商品スペックリスト202309": "aiful",
    "アイフル_除外KW": "aiful",
    "アコム_レギュレーション": "acom",
}


def main():
    xlsx = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_XLSX
    if not xlsx.exists():
        print(f"xlsxが見つからない: {xlsx}", file=sys.stderr)
        return 1
    # 原本保全
    wbdir = RAW / "_workbook"
    wbdir.mkdir(parents=True, exist_ok=True)
    dest = wbdir / f"{date.today():%Y%m%d}_{xlsx.name}"
    if not dest.exists():
        shutil.copy2(xlsx, dest)

    wb = openpyxl.load_workbook(xlsx, data_only=True)
    written = {}
    # 冪等: 既存の wb_*.csv を全slugで削除
    for slug in set(SHEET_MAP.values()):
        d = RAW / slug
        d.mkdir(parents=True, exist_ok=True)
        for old in d.glob("wb_*.csv"):
            old.unlink()
    for sheet, slug in SHEET_MAP.items():
        if sheet not in wb.sheetnames:
            print(f"  [warn] シート無し: {sheet}", file=sys.stderr)
            continue
        ws = wb[sheet]
        out = RAW / slug / f"wb_{sheet.replace(' ', '_').replace('/', '_')}.csv"
        rows = 0
        with open(out, "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f)
            for r in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in r]
                # 末尾の空セルを落とす（巨大シートの空列対策）。
                # 空行もそのまま書く: 行番号を原本シートと一致させ、_labels.json（行番号キーの
                # LLM推定ラベル）を引き続き有効にするため
                while cells and cells[-1] == "":
                    cells.pop()
                w.writerow(cells)
                rows += 1
        written[out.name] = (slug, rows)
    for name, (slug, rows) in written.items():
        print(f"  {slug:12} {name:48} {rows}行")
    print(f"取込完了: {len(written)}シート → regulations/raw/ / 原本={dest.name}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
